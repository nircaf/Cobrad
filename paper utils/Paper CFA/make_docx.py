#!/usr/bin/env python3
"""Build a Word (.docx) version of the paper from the *exact same* prose and
numbers as make_pdf.py: runs make_pdf.py (which rebuilds the PDF as a side
effect, keeping both outputs in sync from one source of truth), captures its
`STORY_FOR_EXPORT` list of reportlab flowables, and re-renders each one as a
native docx paragraph/heading/image -- no PDF text extraction or re-typing of
any sentence.

Run: venv/bin/python "Paper CFA/make_docx.py"
"""
import html
import os
import re
import runpy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.shared import Inches, Pt
from reportlab.platypus import (
    HRFlowable, Image, KeepTogether, PageBreak, Paragraph, Spacer, Table,
)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(HERE, "Cafri_CFA_EEG_paper_v2.docx")

# Map reportlab ParagraphStyle name -> docx rendering.
STYLE_MAP = {
    "PaperTitle": {"heading": 0, "align": "center", "bold": True, "size": 20},
    "Author": {"align": "center", "bold": True, "size": 12},
    "Affil": {"align": "center", "size": 9, "italic": False},
    "AffilList": {"align": "center", "size": 8},
    "H1": {"heading": 1},
    "H2": {"heading": 2},
    "Body": {"align": "justify", "size": 11, "space": (0, 6)},
    "Caption": {"size": 9, "italic": False, "space": (4, 12)},
    "Kw": {"italic": True, "size": 9, "space": (6, 6)},
    "Ref": {"size": 9, "space": (0, 4)},
}

TAG_RE = re.compile(r"<(/?)(super|b|br\s*/?|font[^>]*)>")


def add_markup_runs(paragraph, raw_text):
    """Parse the small subset of reportlab inline markup used in make_pdf.py
    (<super>, <br/>, <font face='Courier'>, HTML entities) into docx runs."""
    pos = 0
    superscript = False
    monospace = False
    bold = False
    for m in TAG_RE.finditer(raw_text):
        chunk = raw_text[pos:m.start()]
        if chunk:
            run = paragraph.add_run(html.unescape(chunk))
            run.font.superscript = superscript
            run.bold = bold or None
            if monospace:
                run.font.name = "Courier New"
        closing, tag = m.group(1), m.group(2)
        if tag == "super":
            superscript = not closing
        elif tag == "b":
            bold = not closing
        elif tag.startswith("br"):
            paragraph.add_run().add_break(WD_BREAK.LINE)
        elif tag.startswith("font"):
            monospace = not closing
        pos = m.end()
    tail = raw_text[pos:]
    if tail:
        run = paragraph.add_run(html.unescape(tail))
        run.font.superscript = superscript
        if monospace:
            run.font.name = "Courier New"


def render_paragraph(doc, flowable):
    style_name = flowable.style.name if flowable.style else "Body"
    spec = STYLE_MAP.get(style_name, {"size": 11})
    if "heading" in spec:
        p = doc.add_heading(level=spec["heading"]) if spec["heading"] else doc.add_paragraph()
    else:
        p = doc.add_paragraph()
    align = spec.get("align")
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_markup_runs(p, flowable.text)
    for run in p.runs:
        if spec.get("bold"):
            run.bold = True
        if spec.get("italic"):
            run.italic = True
        if "size" in spec and "heading" not in spec:
            run.font.size = Pt(spec["size"])
    if "heading" not in spec:
        # exact line height: otherwise Word grows lines that hold a superscript
        # citation, so reference-dense paragraphs look unevenly spaced
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(spec.get("size", 11) * 1.35)
        before, after = spec.get("space", (0, 6))
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
    if style_name == "Ref":
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
    if "heading" in spec and spec["heading"] == 0:
        for run in p.runs:
            run.font.size = Pt(spec["size"])
    return p


def render_image(doc, flowable):
    width_in = flowable.drawWidth / 72.0
    doc.add_picture(flowable.filename, width=Inches(min(width_in, 6.5)))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True


def render_table(doc, flowable):
    rows = flowable._cellvalues
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            table.cell(r, c).width = Inches(flowable._colWidths[c] / 72.0)
            p = table.cell(r, c).paragraphs[0]
            if isinstance(cell, (list, tuple)):  # reportlab wraps laid-out cells
                cell = cell[0]
            add_markup_runs(p, cell.text if isinstance(cell, Paragraph) else str(cell))
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = r == 0


def walk(doc, flowables):
    for f in flowables:
        if isinstance(f, KeepTogether):
            walk(doc, f._content)
        elif isinstance(f, Paragraph):
            render_paragraph(doc, f)
        elif isinstance(f, Image):
            render_image(doc, f)
        elif isinstance(f, Table):
            render_table(doc, f)
        elif isinstance(f, PageBreak):
            doc.add_page_break()
        elif isinstance(f, HRFlowable):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
        elif isinstance(f, Spacer):
            pass
        else:
            pass  # unhandled flowable type, skip silently


print("Running make_pdf.py to get the current story (also rebuilds the PDF)...")
ns = runpy.run_path(os.path.join(HERE, "make_pdf.py"))
story = ns["STORY_FOR_EXPORT"]

doc = Document()
for section in doc.sections:
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)

walk(doc, story)

doc.save(OUT_DOCX)
print(f"Wrote {OUT_DOCX} ({os.path.getsize(OUT_DOCX):,} bytes)")
